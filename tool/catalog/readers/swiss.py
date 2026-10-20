"""Swiss Food Composition Database V 7.1 (FSVO; commercial use allowed with the source cited): raw_ch, per 100 g
edible portion, English names. Its 'Nutrient codes' sheet maps each component name to EuroFIR. Generic foods go to
foods, its branded sheet is products."""

import openpyxl

from normalize import number
from readers.afcd import header


def codes(wb) -> dict[str, str]:
    out = {}
    for r in wb["Nutrient codes"].iter_rows(min_row=4, values_only=True):
        if r and r[1] and r[3]:
            out[header(r[3])[0]] = str(r[1]).strip()  # 'Fat, total (g)' → 'FAT'
    return out


def sheet(wb, name: str, code: dict[str, str], category: str):
    rows = wb[name].iter_rows(values_only=True)
    head = next(r for r in rows if r and r[0] == "ID")
    cols = [header(c) for c in head]  # each value column is followed by "Derivation of value" and "Source"
    for row in rows:
        if not row or row[0] is None:
            continue
        amounts = [
            ("eurofir", code[cols[i][0]], number(v), cols[i][1])
            for i, v in enumerate(row)
            if i >= 8 and cols[i][0] in code
        ]
        yield {
            "code": str(row[0]),
            "name": str(row[3]).strip(),
            "synonyms": row[4],
            "group": row[5],
            "category": category,
            "amounts": amounts,
        }


def read(data):
    wb = openpyxl.load_workbook(data / "raw_ch/Swiss_food_composition_database.xlsx", read_only=True)
    code = codes(wb)
    yield from sheet(wb, "Generic Foods", code, "food")
    yield from sheet(wb, "Branded foods", code, "branded")
