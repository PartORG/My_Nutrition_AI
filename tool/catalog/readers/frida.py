"""Frida 5.5, Danish Food Composition Database (DTU, CC BY 4.0): raw_frida/frida.xlsx, per 100 g; its Parameter
sheet gives each parameter's EuroFIR component and unit, Data_Normalised one value per row."""

from collections import defaultdict

import openpyxl

from normalize import number


def read(data):
    wb = openpyxl.load_workbook(data / "raw_frida/frida.xlsx", read_only=True)
    params = {}
    for r in wb["Parameter"].iter_rows(min_row=2, values_only=True):
        if r[4] is not None and r[8]:
            params[r[4]] = (r[8], (r[2] or "").replace(" ", ""))  # ParameterID → (EuroFIR code, "g/100g")
    amounts, names = defaultdict(list), {}
    for food, name_da, name, param, *_rest in wb["Data_Normalised"].iter_rows(min_row=2, values_only=True):
        names[food] = (name, name_da)
        if param in params:
            amounts[food].append(("eurofir", params[param][0], number(_rest[3]), params[param][1]))
    for food, (name, name_da) in names.items():
        yield {"code": str(food), "name": (name or name_da).strip(), "name_da": name_da, "amounts": amounts[food]}
