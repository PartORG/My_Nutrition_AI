# /// script
# requires-python = ">=3.12"
# dependencies = ["duckdb", "openpyxl"]
# ///
"""Build the real nutri.db from the raw datasets in nutrition-ai-assistant/data.

    uv run tool/build_real_nutri_db.py ~/Programming/nutrition-ai-assistant/data assets/nutri.db

Reads every file in raw_nutritions/ and raw_recipes/ (byte-identical copies once) and writes
docs/DATA_SOURCES.md: per file, how many rows were kept and why the rest were dropped.
Rules (decided 2026-10-06):
  * numbers the app shows as facts must be plausible: non-authoritative rows must pass an
    energy check (kcal ~ 4 protein + 4 carbs + 9 fat) and range checks;
  * a recipe is kept only if every ingredient line maps to a food and a gram amount (so its
    nutrition and allergen filtering are complete) and it states its servings;
  * where a source publishes per-serving calories, the computed value must agree;
  * nutri.db stays under 300 MB: best recipes first.
Conditions, condition rules and medical docs are not in these folders: the development ones
from db/nutri_dummy.sql are carried over unchanged (foods to avoid re-linked by name).
"""

import ast
import csv
import hashlib
import html
import json
import math
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import ingredients as ing  # noqa: E402
from build_nutri_db import finish  # noqa: E402

csv.field_size_limit(sys.maxsize)
ROOT = Path(__file__).resolve().parent.parent
MAX_DB_BYTES = 300_000_000
NUTRIENTS = ("kcal", "protein_g", "fat_g", "carbs_g", "fiber_g", "sugar_g", "sodium_mg")
REPORT: dict[str, dict] = {}  # file -> {"rows": n, "kept": n, "target": str, "dropped": Counter, "note": str}


def report(file: str, target: str, note: str = "") -> dict:
    return REPORT.setdefault(file, {"rows": 0, "kept": 0, "target": target, "dropped": Counter(), "note": note})


def stable_id(*parts) -> int:
    """Ids never change between builds: users' logs and saved recipes point at them (52 bits)."""
    return int(hashlib.sha1(":".join(map(str, parts)).encode()).hexdigest()[:13], 16)


def num(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if isinstance(v, float) and math.isnan(v) else float(v)
    s = str(v).strip().replace(",", ".") if str(v).count(",") == 1 and "." not in str(v) else str(v).strip().replace(",", "")
    m = re.match(r"^[<~]?\s*(-?\d+(?:\.\d+)?(?:e-?\d+)?)", s)
    return float(m[1]) if m else None


def rows(path: Path, **kw):
    with open(path, newline="", encoding=kw.pop("encoding", "utf-8-sig"), errors="replace") as f:
        yield from csv.DictReader(f, **kw)


# --------------------------------------------------------------------------- quality


ALCOHOL = re.compile(r"\b(beer|wine|vodka|whisk|rum|gin|brandy|liqueur|sake|cider|ale|lager|cocktail|"
                     r"margarita|sangria|champagne|martini|bourbon|tequila|alcohol|daiquiri|liquor)", re.I)


def check(n: dict, name: str) -> str | None:
    """None = plausible, else the drop reason."""
    k, p, f, c = n.get("kcal"), n.get("protein_g"), n.get("fat_g"), n.get("carbs_g")
    if None in (k, p, f, c):
        return "missing kcal/protein/fat/carbs"
    if not (0 <= k <= 902) or min(p, f, c) < 0 or max(p, f, c) > 100:
        return "value out of range"
    if p + f + c > 105:
        return "protein+fat+carbs over 100 g per 100 g"
    for key, hi in (("fiber_g", 100), ("sugar_g", 100), ("sodium_mg", 40000)):
        if (v := n.get(key)) is not None and not 0 <= v <= hi:
            return "value out of range"
    if (s := n.get("sugar_g")) is not None and s > c + 1:
        return "sugar above carbs"
    est = 4 * p + 4 * c + 9 * f
    fib = n.get("fiber_g") or 0
    best = min(abs(k - est), abs(k - (est - 2 * min(fib, c))))
    if best > max(20, 0.2 * k):
        if ALCOHOL.search(name) and k > est:  # alcohol adds 7 kcal/g the columns don't show
            return None
        return "kcal does not match protein/fat/carbs"
    return None


def per100(vals: dict, grams: float | None) -> dict | None:
    if not grams or grams <= 0:
        return None
    return {k: (v * 100 / grams if v is not None else None) for k, v in vals.items()}


def tidy(name: str) -> str:
    name = re.sub(r"\s+", " ", (name or "").replace(" ", " ")).strip(" ,;")
    if name.isupper() and len(name) > 3:
        name = name.capitalize()
    return name


def key(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", name.lower()).strip()


# --------------------------------------------------------------------------- foods

SR_ABBREV = [(r"\bW/O\b", "without"), (r"\bWO/", "without "), (r"\bW/", "with "), (r"\bCKD\b", "cooked"),
             (r"\bUNCKD\b", "uncooked"), (r"\bRTE\b", "ready-to-eat"), (r"\bFRZ\b", "frozen"), (r"\bCND\b", "canned"),
             (r"\bDRND\b", "drained"), (r"\bBNLESS\b", "boneless"), (r"\bSKNLESS\b", "skinless"), (r"\bCHS\b", "cheese"),
             (r"\bVEG\b", "vegetable"), (r"\bUNSWTND\b", "unsweetened"), (r"\bSWTND\b", "sweetened"),
             (r"\bPREP\b", "prepared"), (r"\bCOMMLY\b", "commercially"), (r"\bCOMM\b", "commercial"),
             (r"\bNFS\b", "not further specified"), (r"\bENR\b", "enriched"), (r"\bUNENR\b", "unenriched"),
             (r"\bLN\b", "lean"), (r"\bFAT\b", "fat"), (r"\bBRSTD\b", "braised"), (r"\bRSTD\b", "roasted"),
             (r"\bBLD\b", "boiled"), (r"\bSTMD\b", "steamed"), (r"\bWHL\b", "whole"), (r"\bFRSH\b", "fresh"),
             (r"\bPAST\b", "pasteurized"), (r"\bSWT\b", "sweet"), (r"\bCONC\b", "concentrate"), (r"\bH2O\b", "water"),
             (r"\bFLAV\b", "flavored"), (r"\bRED NA\b", "reduced sodium"), (r"\bLO NA\b", "low sodium"),
             (r"\bNA\b", "sodium"), (r"\bAL\b", "aluminum"), (r"\bRTS\b", "ready-to-serve"), (r"\bRTB\b", "ready-to-bake"),
             (r"\bREFR\b", "refrigerated"), (r"\bMXD SP\b", "mixed species"), (r"\bCA\b", "calcium"), (r"\bBNS\b", "beans"), (r"\bSAU\b", "sauce")]


def expand_sr(desc: str) -> str:
    s = desc
    for pat, rep in SR_ABBREV:
        s = re.sub(pat, rep, s)
    return s.replace(",", ", ").lower()


def toks(s: str) -> set[str]:
    return {w.rstrip("s") for w in re.findall(r"[a-z]+", s.lower()) if len(w) > 1}


OPPOSITES = [("with", "without"), ("salted", "unsalted"), ("sweetened", "unsweetened"), ("raw", "cooked"),
             ("enriched", "unenriched"), ("whole", "skim"), ("fresh", "frozen"), ("fresh", "canned")]


def contradicts(a: set[str], b: set[str]) -> bool:
    """'Butter, with salt' vs 'Butter, without salt': one says x where the other says its opposite."""
    return any((x in a and y in b and y not in a) or (y in a and x in b and x not in a) for x, y in OPPOSITES)


def _abbrev(s: str, w: str) -> bool:
    """'soybn' abbreviates 'soybean': same first letter, letters in order."""
    if s == w:
        return True
    if not s or not w or s[0] != w[0] or len(s) > len(w):
        return False
    it = iter(w)
    return all(ch in it for ch in s)


def sr_long_name(short: str, cands: list[tuple]) -> str | None:
    """The long name whose words cover every abbreviated word, with the fewest extra words.
    Several equally good ones (all oils are 884 kcal / 100 g fat) -> None: never guess a name."""
    st = toks(short)
    scored = []
    for name, _, _, lt, *_ in cands:
        missing = sum(1 for s in st if not any(_abbrev(s, w) for w in lt))
        # one word may differ on longer names ('vit c' vs 'ascorbic acid'); nutrients already agree
        if contradicts(st, lt) or missing > (1 if len(st) >= 4 else 0):
            continue
        extra = sum(1 for w in lt if not any(_abbrev(s, w) for s in st))
        scored.append((extra + 2 * missing, len(name), name, frozenset(lt)))
    scored.sort()
    # A tie between two spellings of the same words (both name files) is no tie.
    if not scored or (len(scored) > 1 and scored[0][0] == scored[1][0] and scored[0][3] != scored[1][3]):
        return None
    return scored[0][2]


def load_sr(raw: Path) -> list[dict]:
    """USDA SR28 (ABBREV.csv) with the long names from 'nutrition (2).csv' (matched on
    nutrients + words) and household weights for unit conversion."""
    rn = raw / "raw_nutritions"
    longs = defaultdict(list)
    r2 = report("raw_nutritions/nutrition (2).csv", "foods (names for USDA SR28)",
                "USDA SR28 with long names; used to name ABBREV.csv rows. Same data as nutrition.xlsx.")
    for r in rows(rn / "nutrition (2).csv"):
        r2["rows"] += 1
        kc = num(r["calories"])
        if kc is None:
            r2["dropped"]["no calories"] += 1
            continue
        # This export lists the descriptive parts backwards ('Squash, with salt, ..., zucchini, summer').
        first, *parts = re.split(r",\s*(?![^()]*\))", r["name"])  # commas inside (...) belong to the part
        name = ", ".join([first, *reversed(parts)])
        longs[round(kc)].append((name, num(r["protein"]) or 0, num(r["total_fat"]) or 0, toks(name),
                                 num(r["carbohydrate"]), num(r["sodium"])))
    # FDC SR Legacy carries the same foods with their full names: a second chance at a name.
    for r in rows(rn / "comprehensive_foods_usda.csv"):
        if r["data_type"] == "SR Legacy" and (kc := num(r["calories"])) is not None:
            name = tidy(r["food_name"])
            longs[round(kc)].append((name, num(r["protein_g"]) or 0, num(r["fat_g"]) or 0, toks(name),
                                     num(r["carbs_g"]), num(r["sodium_mg"])))
    out, named = [], 0
    rep = report("raw_nutritions/ABBREV.csv", "foods (generic, authoritative)", "USDA SR28 abbreviated file.")
    for r in rows(rn / "ABBREV.csv"):
        rep["rows"] += 1
        n = {"kcal": num(r["Energ_Kcal"]), "protein_g": num(r["Protein_(g)"]), "fat_g": num(r["Lipid_Tot_(g)"]),
             "carbs_g": num(r["Carbohydrt_(g)"]), "fiber_g": num(r["Fiber_TD_(g)"]), "sugar_g": num(r["Sugar_Tot_(g)"]),
             "sodium_mg": num(r["Sodium_(mg)"])}
        if None in (n["kcal"], n["protein_g"], n["fat_g"], n["carbs_g"]):
            rep["dropped"]["missing kcal/protein/fat/carbs"] += 1
            continue
        short = expand_sr(r["Shrt_Desc"])
        name = sr_long_name(short, [c for c in longs.get(round(n["kcal"]), [])
                                    if abs(c[1] - n["protein_g"]) < 0.6 and abs(c[2] - n["fat_g"]) < 1.1
                                    and (c[4] is None or abs(c[4] - n["carbs_g"]) < 1.1)
                                    and (c[5] is None or n["sodium_mg"] is None
                                         or abs(c[5] - n["sodium_mg"]) <= max(2, 0.05 * n["sodium_mg"]))])
        named += name is not None
        name = name or short.capitalize()
        weights = []
        for g, d in ((r["GmWt_1"], r["GmWt_Desc1"]), (r["GmWt_2"], r["GmWt_Desc2"])):
            if (gv := num(g)) and d:
                weights.append((d.lower(), gv))
        out.append({"name": tidy(name), "src": "usda_sr28", "skey": r["NDB_No"], "generic": True, **n,
                    "serving_g": weights[0][1] if weights else None, "weights": weights,
                    "refuse": num(r["Refuse_Pct"]) or 0})
        rep["kept"] += 1
    r2["kept"] = named
    r2["note"] += f" Named {named} of {len(out)} SR28 foods; the rest keep the expanded short name."
    r2["dropped"]["no unambiguous ABBREV.csv match (name not used)"] = r2["rows"] - named
    rep_x = report("raw_nutritions/nutrition.xlsx", "same as nutrition (2).csv",
                   "Same 8789 USDA SR28 rows and names as nutrition (2).csv (checked row by row); not loaded twice.")
    rep_x["rows"] = r2["rows"]
    rep_x["dropped"]["duplicate of nutrition (2).csv"] = r2["rows"]
    return out


def load_food_csv(raw: Path, sr: list[dict]) -> list[dict]:
    """food.csv is USDA SR (7413 foods, same NDB numbers); adds household weights SR lacks."""
    by_ndb = {f["skey"]: f for f in sr}
    rep = report("raw_nutritions/food.csv", "foods (merged into USDA SR28)", "USDA SR, same NDB numbers as ABBREV.csv.")
    out = []
    for r in rows(raw / "raw_nutritions/food.csv"):
        rep["rows"] += 1
        ndb = r["Nutrient Data Bank Number"]
        if f := by_ndb.get(ndb):
            for g, d in ((r["Data.Household Weights.1st Household Weight"], r["Data.Household Weights.1st Household Weight Description"]),
                         (r["Data.Household Weights.2nd Household Weight"], r["Data.Household Weights.2nd Household Weight Description"])):
                if (gv := num(g)) and d and (d.lower(), gv) not in f["weights"]:
                    f["weights"].append((d.lower(), gv))
            rep["dropped"]["same food as ABBREV.csv (merged)"] += 1
            continue
        n = {"kcal": num(r["Data.Kilocalories"]), "protein_g": num(r["Data.Protein"]), "fat_g": num(r["Data.Fat.Total Lipid"]),
             "carbs_g": num(r["Data.Carbohydrate"]), "fiber_g": num(r["Data.Fiber"]), "sugar_g": num(r["Data.Sugar Total"]),
             "sodium_mg": num(r.get("Data.Major Minerals.Sodium"))}
        if None in (n["kcal"], n["protein_g"], n["fat_g"], n["carbs_g"]):
            rep["dropped"]["missing kcal/protein/fat/carbs"] += 1
            continue
        out.append({"name": tidy(expand_sr(r["Description"]).capitalize()), "src": "usda_sr", "skey": ndb, "generic": True, **n,
                    "serving_g": num(r["Data.Household Weights.1st Household Weight"]), "weights": []})
        rep["kept"] += 1
    return out


def load_fdc(raw: Path) -> list[dict]:
    """comprehensive_foods_usda.csv: FDC SR Legacy (same as SR28), Foundation, Branded."""
    rep = report("raw_nutritions/comprehensive_foods_usda.csv", "foods (Foundation generic, Branded)",
                 "USDA FoodData Central extract, per 100 g.")
    out = []
    for r in rows(raw / "raw_nutritions/comprehensive_foods_usda.csv"):
        rep["rows"] += 1
        t = r["data_type"]
        if t == "SR Legacy":
            rep["dropped"]["SR Legacy = same foods as ABBREV.csv"] += 1
            continue
        name = tidy(r["food_name"])
        if t == "Branded" and r["brand_name"] and r["brand_name"].lower() not in name.lower():
            name = f"{name} ({tidy(r['brand_name']).title()})"
        n = {"kcal": num(r["calories"]), "protein_g": num(r["protein_g"]), "fat_g": num(r["fat_g"]),
             "carbs_g": num(r["carbs_g"]), "fiber_g": num(r["fiber_g"]), "sugar_g": num(r["sugar_g"]),
             "sodium_mg": num(r["sodium_mg"])}
        if t == "Branded" and (why := check(n, name)):
            rep["dropped"][why] += 1
            continue
        if None in (n["kcal"], n["protein_g"], n["fat_g"], n["carbs_g"]):
            rep["dropped"]["missing kcal/protein/fat/carbs"] += 1
            continue
        serving = num(r["serving_size"]) if r["serving_unit"] in ("GRM", "g") else None
        out.append({"name": name, "src": "usda_fdc_" + t.lower(), "skey": r["fdc_id"], "generic": t == "Foundation", **n,
                    "serving_g": serving, "weights": [], "ingredients": r["ingredients"]})
        rep["kept"] += 1
    return out


def load_nz(raw: Path) -> list[dict]:
    import openpyxl
    f = raw / "raw_nutritions/Nutritional Datasets for New Zealand Foods.xlsx"
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    name = "raw_nutritions/" + f.name
    rep = report(name, "foods (generic, authoritative)",
                 "NZ FOODfiles 2018. Sheet FOODFiles-Nutrients loaded. FOODFiles-AminoAcids: amino acids only (not in the "
                 "schema). Matching Table: formulas pointing back at FOODFiles-Nutrients (duplicates). 'Recipes - Drew et al "
                 "2020': ingredient percentages for survey food groups, no instructions or amounts (cannot be recipes). "
                 "Introduction: text.")
    out = []
    ws = wb["FOODFiles-Nutrients"]
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0] or not r[2]:
            continue
        rep["rows"] += 1
        n = {"kcal": num(r[4]), "protein_g": num(r[9]), "fat_g": num(r[7]), "carbs_g": num(r[58]),
             "fiber_g": num(r[10]), "sugar_g": num(r[5]), "sodium_mg": num(r[33])}
        if None in (n["kcal"], n["protein_g"], n["fat_g"], n["carbs_g"]):
            rep["dropped"]["missing kcal/protein/fat/carbs"] += 1
            continue
        out.append({"name": tidy(str(r[2])), "src": "nz_foodfiles", "skey": r[0], "generic": True, **n,
                    "serving_g": None, "weights": []})
        rep["kept"] += 1
    for sheet in ("FOODFiles-AminoAcids", "Matching Table ", "Recipes - Drew et al 2020", "Introduction"):
        ws = wb[sheet]
        rep["dropped"][f"sheet '{sheet.strip()}' not nutrition per food (see note)"] += max(0, (ws.max_row or 0) - 5)
    return out


def simple(raw: Path, file: str, target: str, note: str, src: str, name_col, cols: dict, *, grams=None,
           generic=False, encoding="utf-8-sig", skip=None, scale=None, serving=None, delimiter=",") -> list[dict]:
    """One flat per-100 g (or per-serving with known grams) CSV -> checked food rows."""
    rep = report("raw_nutritions/" + file, target, note)
    out = []
    for i, r in enumerate(rows(raw / "raw_nutritions" / file, encoding=encoding, delimiter=delimiter)):
        rep["rows"] += 1
        name = tidy(name_col(r) if callable(name_col) else r[name_col])
        if not name:
            rep["dropped"]["no name"] += 1
            continue
        if skip and (why := skip(r)):
            rep["dropped"][why] += 1
            continue
        n = {k: (c(r) if callable(c) else num(r.get(c))) if c else None for k, c in cols.items()}
        if scale:
            n = scale(n)
        if grams:
            g = grams(r)
            n = per100(n, g)
            if n is None:
                rep["dropped"]["no gram weight for the serving"] += 1
                continue
        if why := check(n, name):
            rep["dropped"][why] += 1
            continue
        out.append({"name": name, "src": src, "skey": i, "generic": generic, **n,
                    "serving_g": (grams(r) if grams else serving(r) if serving else None), "weights": []})
        rep["kept"] += 1
    return out


def grams_in(text: str) -> float | None:
    """'4.8 oz (136 g)' -> 136; '143.5 g' -> 143.5; '(93 g/3.3 oz)' -> 93; '8 oz' -> 226.8."""
    if not text:
        return None
    if m := re.search(r"(\d+(?:\.\d+)?)\s*(?:g|gm|grams?)\b", text, re.I):
        return float(m[1])
    if m := re.search(r"(\d+(?:\.\d+)?)\s*(?:oz|ounces?)\b", text, re.I):
        return float(m[1]) * 28.35
    if m := re.search(r"(\d+(?:\.\d+)?)\s*(?:ml)\b", text, re.I):
        return float(m[1])
    return None


def load_flat_foods(raw: Path) -> tuple[list[dict], list[dict]]:
    """All the smaller files. Returns (foods, dishes)."""
    foods, dishes = [], []
    std = {"kcal": "calories", "protein_g": "protein_g", "fat_g": "fat_g", "carbs_g": "carbs_g",
           "fiber_g": "fiber_g", "sugar_g": "sugar_g", "sodium_mg": "sodium_mg"}
    foods += simple(raw, "healthy_foods_database.csv", "foods", "USDA-style names, per 100 g.", "healthy_foods_db",
                    "food_name", std)
    foods += simple(raw, "Food_Nutrition_Dataset.csv", "foods", "USDA-style names, per 100 g; no fiber/sugar/sodium.",
                    "food_nutrition_dataset", "food_name",
                    {"kcal": "calories", "protein_g": "protein", "fat_g": "fat", "carbs_g": "carbs"})
    foods += simple(raw, "foodstruct_nutritional_facts.csv", "foods (generic)", "foodstruct.com, per 100 g; minerals in grams.",
                    "foodstruct", "Food Name",
                    {"kcal": "Calories", "protein_g": "Protein", "fat_g": "Fats", "carbs_g": "Carbs", "fiber_g": "Fiber",
                     "sugar_g": "Sugar", "sodium_mg": lambda r: (num(r["Sodium"]) * 1000 if num(r["Sodium"]) is not None else None)},
                    generic=True)
    foods += simple(raw, "Indian_Food_Nutrition_Processed.csv", "foods (Indian dishes, per 100 g)",
                    "Indian Nutrient Databank style, per 100 g; no serving weight so they are foods, not camera dishes.",
                    "indb", "Dish Name",
                    {"kcal": "Calories (kcal)", "protein_g": "Protein (g)", "fat_g": "Fats (g)", "carbs_g": "Carbohydrates (g)",
                     "fiber_g": "Fibre (g)", "sugar_g": "Free Sugar (g)", "sodium_mg": "Sodium (mg)"})
    foods += simple(raw, "nutrition.csv", "foods (Indonesian foods, per 100 g)", "Indonesian food table, per 100 g.",
                    "indonesia", "name",
                    {"kcal": "calories", "protein_g": "proteins", "fat_g": "fat", "carbs_g": "carbohydrate"})
    foods += simple(raw, "nutrients_csvfile.csv", "foods", "Per household measure with grams; converted to per 100 g.",
                    "nutrients_csvfile", "Food",
                    {"kcal": "Calories", "protein_g": "Protein", "fat_g": "Fat", "carbs_g": "Carbs", "fiber_g": "Fiber"},
                    grams=lambda r: num(r["Grams"]),
                    skip=lambda r: "section header row" if not num(r["Grams"]) and not num(r["Calories"]) else None)

    def baked_fat(r):
        parts = [num(r[c]) for c in ("SaturatedFat-g", "MonounsaturatedFat-g", "PolyunsaturatedFat-g", "TransaFat-g")]
        return sum(p for p in parts if p is not None) if any(p is not None for p in parts) else None

    foods += simple(raw, "BakedFoodNutritions.csv", "foods (branded baked goods)",
                    "Per 100 g; total fat = sum of the fatty-acid columns (no total fat column); 'Sodium-g' is in mg.",
                    "baked_foods", lambda r: f"{r['ProductType']}, {r['FlavorVariant'].strip()} ({r['Food']})".replace(", ()", ""),
                    {"kcal": "Calories-kcl", "protein_g": "Protein-g", "fat_g": baked_fat, "carbs_g": "Carb-g",
                     "fiber_g": "Fiber-g", "sugar_g": "Sugar-g", "sodium_mg": "Sodium-g"},
                    serving=lambda r: num(r["MQGmWt"]))

    # OFF extracts: product names + per 100 g, no barcode -> branded foods (allergen flags kept).
    off = simple(raw, "foods_health_scores_allergens.csv", "foods (branded, with allergens)",
                 "Open Food Facts extract without barcodes, per 100 g; allergen flags become allergens rows.",
                 "off_extract", lambda r: (f"{r['product_name']} ({r['brands'].split(',')[0].strip()})" if r["brands"] else r["product_name"]),
                 {"kcal": "energy_kcal", "protein_g": "proteins_100g", "fat_g": "fat_100g", "carbs_g": "carbs_100g",
                  "fiber_g": "fiber_100g", "sugar_g": "sugars_100g",
                  "sodium_mg": lambda r: num(r["sodium_100g"]) * 1000 if num(r["sodium_100g"]) is not None else None})
    # Allergen flags from all three OFF extracts, by name+brand; the two without nutrition only add flags.
    flags: dict[str, set[str]] = defaultdict(set)
    flag_cols = (("contains_gluten", "gluten"), ("contains_dairy", "milk"), ("contains_nuts", "tree nuts"),
                 ("contains_soy", "soy"), ("contains_eggs", "egg"), ("contains_fish", "fish"))
    hs = list(rows(raw / "raw_nutritions/foods_health_scores_allergens.csv"))
    pkey = lambda r: key(r["product_name"] + " " + (r["brands"] or ""))  # noqa: E731
    hs_keys = {pkey(x) for x in hs}
    for file in ("foods_health_scores_allergens.csv", "foods_allergens.csv", "foods_dietary_restrictions.csv"):
        file_rows = hs if file == "foods_health_scores_allergens.csv" else list(rows(raw / "raw_nutritions" / file))
        for r in file_rows:
            flags[pkey(r)] |= {a for col, a in flag_cols if r.get(col) == "True"} | allergens_from_text(r.get("allergens") or "")
        if file != "foods_health_scores_allergens.csv":
            rep = report("raw_nutritions/" + file, "allergen flags for foods_health_scores_allergens.csv",
                         "Open Food Facts extract: names and allergen flags, no nutrition values.")
            rep["rows"] = len(file_rows)
            matched = sum(1 for r in file_rows if pkey(r) in hs_keys)
            rep["kept"] = matched
            rep["dropped"]["no nutrition values and no matching product with nutrition"] = rep["rows"] - matched
            rep["note"] += f" {matched} rows add allergen flags to a product of foods_health_scores_allergens.csv."
    for f in off:  # skey = row index in the file
        f["allergen_flags"] = flags[pkey(hs[f["skey"]])]
    foods += off

    # Fruit/veg/seafood FDA posters: grams in the serving text.
    for file, enc in (("nutritionalfacts_fruit_vegetables_seafood.csv", "latin-1"),):
        foods += simple(raw, file, "foods (generic)", "FDA raw fruit/vegetable/seafood labels; serving grams in the name.",
                        "fda_produce", lambda r: re.split(r",|\(", r["food_and_serving"])[0],
                        {"kcal": "calories", "protein_g": "protein_g", "fat_g": "total_fat_g", "carbs_g": "total_carbo_hydrate_g",
                         "fiber_g": "dietary_fiber_g", "sugar_g": "sugars_g",
                         "sodium_mg": "sodium_g"},  # the column says g but holds mg (asparagus: 0)
                        grams=lambda r: grams_in(r["food_and_serving"]) or grams_in(r["food_type"]),
                        generic=True, encoding=enc)
    rep = report("raw_nutritions/NutritionalFacts_Fruit_Vegetables_Seafood.csv", "same as the lower-case file",
                 "Older export of nutritionalfacts_fruit_vegetables_seafood.csv (no header row, 54 of its 61 rows).")
    for _ in rows(raw / "raw_nutritions/NutritionalFacts_Fruit_Vegetables_Seafood.csv", encoding="latin-1"):
        rep["rows"] += 1
    rep["dropped"]["duplicate of nutritionalfacts_fruit_vegetables_seafood.csv"] = rep["rows"]

    # ---- dishes (restaurant / prepared, with a serving weight)
    def dish_file(file, note, src, name_col, cols, grams, delimiter=","):
        dishes.extend(simple(raw, file, "dishes", note, src, name_col, cols, grams=grams, delimiter=delimiter))

    dish_file("menu.csv", "McDonald's menu; serving weight in 'Serving Size'.", "mcdonalds_menu",
              lambda r: f"{r['Item']} (McDonald's)",
              {"kcal": "Calories", "protein_g": "Protein", "fat_g": "Total Fat", "carbs_g": "Carbohydrates",
               "fiber_g": "Dietary Fiber", "sugar_g": "Sugars", "sodium_mg": "Sodium"},
              lambda r: grams_in(r["Serving Size"]))
    dish_file("nutrition_values.csv", "Fast-food chains with serving grams.", "chains_nutrition_values",
              lambda r: f"{r['Item']} ({r['Chain']})",
              {"kcal": "Calories", "protein_g": "Protein (g)", "fat_g": "Total Fat (g)", "carbs_g": "Total Carb (g)",
               "fiber_g": "Dietary Fiber (g)", "sugar_g": "Total Sugar (g)", "sodium_mg": "Sodium (mg)"},
              lambda r: num(r["Serving Size (g)"]), delimiter=";")
    dish_file("Nutrition_Value_Dataset.csv", "Indian fast-food chains, 'Per Serve Size' in g.", "india_chains",
              lambda r: f"{r['Product']} ({r['Company']})",
              {"kcal": "Energy (kCal)", "protein_g": "Protein (g)", "fat_g": "Total Fat (g)", "carbs_g": "Carbohydrates (g)",
               "fiber_g": "Fiber (g)", "sugar_g": "Sugar (g)", "sodium_mg": "Sodium (mg)"},
              lambda r: grams_in(r["Per Serve Size"]))
    dish_file("daily_food_nutrition_dataset.csv", "Per serving; only rows whose name states grams/ounces can be used.",
              "daily_food", "Food_Item",
              {"kcal": "Calories (kcal)", "protein_g": "Protein (g)", "fat_g": "Fat (g)", "carbs_g": "Carbohydrates (g)",
               "fiber_g": "Fiber (g)", "sugar_g": "Sugars (g)", "sodium_mg": "Sodium (mg)"},
              lambda r: grams_in(r["Food_Item"]))

    # Food-101 labels x 5 portion sizes: same per-100 g, one dish per label (median portion).
    rep = report("raw_nutritions/nutrition (1).csv", "dishes (Food-101 labels)",
                 "101 Food-101 dish labels at 5 portion weights; one dish per label at the middle portion.")
    byl = defaultdict(list)
    for r in rows(raw / "raw_nutritions/nutrition (1).csv"):
        rep["rows"] += 1
        byl[r["label"]].append(r)
    for label, rs in byl.items():
        rs.sort(key=lambda r: num(r["weight"]) or 0)
        r = rs[len(rs) // 2]
        name = label.replace("_", " ")
        n = per100({"kcal": num(r["calories"]), "protein_g": num(r["protein"]), "fat_g": num(r["fats"]),
                    "carbs_g": num(r["carbohydrates"]), "fiber_g": num(r["fiber"]), "sugar_g": num(r["sugars"]),
                    "sodium_mg": num(r["sodium"])}, num(r["weight"]))
        if why := check(n, name):
            rep["dropped"][why] += len(rs)
            continue
        dishes.append({"name": name, "src": "food101", "skey": label, "generic": False, **n, "serving_g": num(r["weight"]), "weights": []})
        rep["kept"] += 1
        rep["dropped"]["other portion sizes of the same dish (same per-100 g)"] += len(rs) - 1

    # ---- files that cannot give per-100 g values
    def unusable(file, target, note, reason, enc="utf-8"):
        rep = report("raw_nutritions/" + file, target, note)
        for _ in rows(raw / "raw_nutritions" / file, encoding=enc):
            rep["rows"] += 1
        rep["dropped"][reason] = rep["rows"]

    unusable("fastfood.csv", "-", "Fast-food items per serving.", "no serving weight (per-100 g impossible)")
    unusable("FastFoodNutritionMenuV2.csv", "-", "Fast-food items per serving.", "no serving weight (per-100 g impossible)")
    unusable("FastFoodNutritionMenuV3.csv", "-", "Same as V2 with one row fewer.", "no serving weight (per-100 g impossible)")
    unusable("starbucks.csv", "-", "Starbucks food items per serving.", "no serving weight (per-100 g impossible)")
    unusable("calorie_dataset.csv", "-", "Macro numbers with a calorie class, for ML.", "no food names")
    unusable("food_menu_nutrition_dataset.csv", "-",
             "Generated: random chain/city/price combinations (e.g. a Five Guys 'Veggie Sub' in London).",
             "synthetic data")
    unusable("healthy_eating_dataset.csv", "-", "Generated: random meal names x cuisine x diet ('Kid Pasta, Indian, Keto').",
             "synthetic data")
    unusable("Food Composition.csv", "-",
             "Australian Food Composition (AFCD) food details: descriptions and sampling notes, no nutrient values.",
             "no nutrient values")
    unusable("Top 100 Healthiest Food in the World.csv", "-", "Calories, protein, fiber only.", "no fat/carbs values")
    for g in range(1, 6):
        unusable(f"FOOD-DATA-GROUP{g}.csv", "-",
                 "Per serving with unknown weight (cream cheese 51 kcal = one tablespoon).",
                 "per serving with no serving weight")
    unusable("cleaned_nutrition_dataset_per100g.csv", "-",
             "Claims per 100 g but holds the FOOD-DATA-GROUP per-serving values (margarine with yoghurt: 88 kcal).",
             "per serving with no serving weight")
    return foods, dishes


ALLERGEN_WORDS = [  # (regex on a food name or ingredient line, allergen). Conservative: a false 'contains' is safe.
    (r"\b(peanut|satay|arachis)", "peanuts"),
    (r"\b(almond|walnut|pecan|cashew|pistachio|hazelnut|filbert|macadamia|brazil ?nut|pine ?nut|pignoli|chestnut|"
     r"praline|marzipan|nutella|gianduja|mixed nuts|nut butter|nuts\b|nut\b|frangipane|pesto|baklava|nougat)", "tree nuts"),
    (r"\b(milk|cheese|chees|cheez|butter(?!nut|fl)|cream(?!y|ed\b)|cr[eè]me|yogh?urt|whey|casein|lactose|ghee|kefir|paneer|"
     r"ricotta|mozzarella|parmesan|parmigiano|pecorino|romano|cheddar|gruyere|provolone|custard|ice cream|gelato|"
     r"buttermilk|dairy|curd|quark|mascarpone|brie|camembert|feta|gouda|latte|cappuccino|milkshake|chocolate|fudge|"
     r"caramel|toffee|alfredo|bechamel|b[eé]chamel|pudding|eggnog|raita|lassi|kheer|tzatziki|cheesecake|queso|"
     r"halloumi|labneh|half[- ]and[- ]half|velveeta|whipped topping|cool whip|dulce de leche|nougat|ganache)", "milk"),
    (r"\b(eggs?\b|egg (?:white|yolk|noodle|substitute|roll|wash)|mayonnaise|mayo\b|meringue|aioli|omelet|"
     r"frittata|quiche|custard|eggnog|hollandaise|b[eé]arnaise|souffl|pancake|waffle|brioche|challah|crepe|"
     r"french toast|carbonara|ladyfinger|lady finger)", "egg"),
    (r"\b(wheat|flour|bread|pasta|noodle|spaghett|macaroni|couscous|bulgur|semolina|barley|rye\b|spelt|farro|"
     r"seitan|cracker|biscuit|biscotti|cookie|cake|cupcake|muffin|bagel|croissant|pastry|pie\b|pies\b|pizza|"
     r"tortilla|breadcrumb|crumbs?\b|panko|crouton|cereal|bran\b|malt|"
     r"beer|ale\b|lager|soy sauce|shoyu|teriyaki|hoisin|udon|soba|ramen|gnocchi|dumpling|wonton|wrap\b|pita|"
     r"naan|chapati|roti\b|paratha|bun\b|buns\b|roll\b|rolls\b|donut|doughnut|dough|pretzel|waffle|pancake|"
     r"brownie|granola|oat|lasagn|ravioli|tortellini|penne|fettuccin|linguin|rigatoni|ziti|vermicelli|farfalle|"
     r"rotini|orzo|phyllo|filo\b|matzo|batter\b|breaded|gravy|stuffing|scone|crescent|durum|farina|graham|"
     r"focaccia|ciabatta|baguette|pumpernickel|sponge|ritz|oreo|bisquick|cheerios|gluten|kamut|freekeh|triticale)",
     "gluten"),
    (r"((?<!shell)fish|\b(salmon|tuna|cod\b|haddock|halibut|trout|sardine|anchov|mackerel|herring|tilapia|pollock|"
     r"sole\b|bass\b|snapper|swordfish|carp\b|perch|pike\b|eel\b|caviar|roe\b|worcestershire|caesar|surimi|"
     r"bonito|dashi|flounder|mahi|pangasius|sushi|lox\b|kipper|branzino|grouper|nam pla|bacalao|sprat|whiting|"
     r"hake|turbot|barramundi))", "fish"),
    (r"\b(shrimp|prawn|crab|lobster|crayfish|crawfish|scallop(?!ed)|clam|mussel|oyster(?! mushroom)|squid|calamari|"
     r"octopus|shellfish|langoustine|langostino|scampi|cuttlefish|abalone|whelk|cockle|snail|escargot|seafood|krill)",
     "shellfish"),
    (r"\b(soy|soya|tofu|tempeh|edamame|miso|soybean|tamari|teriyaki|hoisin|natto|bean curd|tvp\b|"
     r"textured vegetable protein|shoyu)", "soy"),
    (r"\b(sesame|tahini|halva|halvah|hummus|gomasio|benne|za'?atar)", "sesame"),
]
_ALLERGEN_RE = [(re.compile(p, re.I), a) for p, a in ALLERGEN_WORDS]
# Harmless phrases rewritten before matching: 'almond milk' keeps its nut, loses the milk.
BENIGN = [(re.compile(p, re.I), r) for p, r in (
    (r"\b(almond|soy|soya|oat|rice|coconut|cashew|hemp|hazelnut|macadamia|pea|flax)[- ]milk", r"\1 drink"),
    (r"\b(peanut|almond|cashew|sunflower|cocoa|apple|shea|nut|seed|pumpkin) butter", r"\1 spread"),
    (r"\bwater ?chestnuts?", "waterchestnut"), (r"\bcoconut", "coco"), (r"\bnutmeg", "spice"),
    (r"\bbutternut", "squash"), (r"\bcream of tartar", "tartar"), (r"\bnutritional yeast", "yeast"),
    (r"\bcorn tortillas?|tortillas?,? (?:ready-to-bake or -fry, )?corn|tortilla chips", "maize flatbread"),
    (r"\bbutter beans?", "lima beans"), (r"\bbutter[- ]flavou?r(ed)?", "flavored"), (r"\bbread machine", "machine"),
    (r"\bmasa( harina)?", "maize"), (r"\b(gluten|dairy|nut|egg|soy)[- ]free\b", "free"), (r"\bnon-?dairy", "plant"), (r"\bvegan (cheese|butter|mayo\w*)", r"plant"),
)]
GLUTEN_FREE_FLOUR = re.compile(r"\b(rice|corn|almond|coconut|buckwheat|potato|tapioca|chickpea|gram|cassava|sorghum|"
                               r"teff|quinoa|millet|arrowroot|soy|amaranth|maize) ?(flour|starch|noodles?|pasta)", re.I)


def allergens_from_text(text: str) -> set[str]:
    """Allergens a food name or ingredient line implies. False positives are acceptable, misses are not."""
    t = text.lower()
    for rx, rep in BENIGN:
        t = rx.sub(rep, t)
    out = {a for rx, a in _ALLERGEN_RE if rx.search(t)}
    if re.search(r"gluten[- ]free", text, re.I):
        out.discard("gluten")
    if "gluten" in out and GLUTEN_FREE_FLOUR.search(t) and not re.search(r"\b(wheat|rye|barley|spelt|bread flour|"
                                                                         r"all-purpose|self-ris)", t):
        out.discard("gluten")
    return out


def load_off_products(raw: Path) -> list[dict]:
    """Open Food Facts 2019 dump -> products (barcode lookup)."""
    f = raw / "raw_nutritions/en.openfoodfacts.org.products.tsv"
    rep = report("raw_nutritions/" + f.name, "products (barcodes)",
                 "Open Food Facts dump (Sep 2019). energy_100g is kJ; sodium_100g in g (or from salt_100g / 2.5).")
    out, seen = [], set()
    with open(f, encoding="utf-8", errors="replace") as fh:
        head = fh.readline().rstrip("\n").split("\t")
        ix = {h: i for i, h in enumerate(head)}
        for line in fh:
            rep["rows"] += 1
            c = line.rstrip("\n").split("\t")
            if len(c) != len(head):
                rep["dropped"]["malformed line"] += 1
                continue
            got = off_product(lambda k: c[ix[k]] if k in ix else "")
            if isinstance(got, str):
                rep["dropped"][got] += 1
                continue
            code = got["barcode"]
            if code in seen:
                rep["dropped"]["duplicate barcode"] += 1
                continue
            seen.add(code)
            out.append(got)
            rep["kept"] += 1
    return out


def off_product(get) -> dict | str:
    """One Open Food Facts row (get(column) -> text) -> a products row, or why it was dropped.
    Shared by the 2019 dump in nutri.db and the current dump in products.db (tool/build_products_db.py)."""
    code = get("code").strip()
    if not re.fullmatch(r"\d{8,14}", code):
        return "no valid barcode"
    name = tidy(get("product_name"))
    if not name:
        return "no product name"
    brand = get("brands").split(",")[0].strip()
    if brand and brand.lower() not in name.lower():
        name = f"{name} ({brand})"
    kcal = num(get("energy-kcal_100g"))
    if kcal is None and (kj := num(get("energy_100g"))) is not None:
        kcal = kj / 4.184
    sod = num(get("sodium_100g"))
    if sod is None and (salt := num(get("salt_100g"))) is not None:
        sod = salt / 2.5
    n = {"kcal": kcal, "protein_g": num(get("proteins_100g")),
         "fat_g": num(get("fat_100g")), "carbs_g": num(get("carbohydrates_100g")),
         "fiber_g": num(get("fiber_100g")), "sugar_g": num(get("sugars_100g")),
         "sodium_mg": sod * 1000 if sod is not None else None}
    if why := check(n, name):
        return why
    return {"barcode": code.zfill(13) if len(code) in (12, 13) else code, "name": name[:200], **n}


# --------------------------------------------------------------------------- foods -> one table

FOOD_PRIORITY = ["usda_sr28", "usda_sr", "usda_fdc_foundation", "nz_foodfiles", "foodstruct", "fda_produce",
                 "healthy_foods_db", "food_nutrition_dataset", "nutrients_csvfile", "indb", "indonesia",
                 "usda_fdc_branded", "baked_foods", "off_extract"]


def merge_foods(groups: list[list[dict]]) -> list[dict]:
    """One row per name, the most authoritative source winning; stable ids; allergens from the name,
    the ingredient list (branded) and the source's own flags."""
    allf = sorted((f for g in groups for f in g), key=lambda f: FOOD_PRIORITY.index(f["src"]))
    out, seen = [], {}
    for f in allf:
        k = key(f["name"])
        if not k:
            continue
        if k in seen:
            report_src(f["src"])["dropped"]["same name as a more authoritative food"] += 1
            report_src(f["src"])["kept"] -= 1
            continue
        f["id"] = stable_id("food", f["src"], f["skey"])
        f["allergens"] = allergens_from_text(f["name"] + " " + (f.get("ingredients") or "")) | f.get("allergen_flags", set())
        seen[k] = f
        out.append(f)
    return out


SRC_FILE = {"usda_sr28": "raw_nutritions/ABBREV.csv", "usda_sr": "raw_nutritions/food.csv",
            "usda_fdc_foundation": "raw_nutritions/comprehensive_foods_usda.csv",
            "usda_fdc_branded": "raw_nutritions/comprehensive_foods_usda.csv",
            "nz_foodfiles": "raw_nutritions/Nutritional Datasets for New Zealand Foods.xlsx",
            "foodstruct": "raw_nutritions/foodstruct_nutritional_facts.csv",
            "fda_produce": "raw_nutritions/nutritionalfacts_fruit_vegetables_seafood.csv",
            "healthy_foods_db": "raw_nutritions/healthy_foods_database.csv",
            "food_nutrition_dataset": "raw_nutritions/Food_Nutrition_Dataset.csv",
            "nutrients_csvfile": "raw_nutritions/nutrients_csvfile.csv", "indb": "raw_nutritions/Indian_Food_Nutrition_Processed.csv",
            "indonesia": "raw_nutritions/nutrition.csv", "baked_foods": "raw_nutritions/BakedFoodNutritions.csv",
            "off_extract": "raw_nutritions/foods_health_scores_allergens.csv"}


def report_src(src: str) -> dict:
    return REPORT[SRC_FILE[src]]


# --------------------------------------------------------------------------- ingredient -> food + grams

VOLUME_DESC = re.compile(r"^\s*([\d.]+(?:\s+\d+/\d+)?|\d+/\d+)?\s*(cups?|tbsp|tablespoons?|tsp|teaspoons?|fl oz|ml|liters?|pints?|quarts?)\b")
LIQUID = [(re.compile(r"^(oil|shortening|lard|butter oil)", re.I), 0.92), (re.compile(r"^(syrups?|molasses|corn syrup)", re.I), 1.37),
          (re.compile(r"^honey", re.I), 1.42),
          (re.compile(r"^(milk|cream, fluid|beverages?|alcoholic|water|soup, .*(broth|stock|ready)|vinegar|.*juice|soy sauce|"
                      r"sauce, (fish|worcestershire|teriyaki|soy|ready-to-serve, pepper)|vanilla extract|coffee|wine)", re.I), 1.0)]
SIZE = {"large": 1.25, "small": 0.75, "medium": 1.0, "whole": 1.0, None: 1.0}


# 'tilapia fillets' is tilapia: the cut is not part of the food's name.
CUT_WORDS = {"fillet", "steak", "chop", "piece", "halve", "half", "floret", "spear", "cube", "slice", "strip",
             "chunk", "wedge", "link"}


class Matcher:
    """Normalised ingredient name -> (food, piece grams, grams per cup, extra allergens)."""

    def __init__(self, foods: list[dict], map_file: Path):
        self.by_name = {f["name"]: f for f in foods}
        self.map: dict[str, tuple] = {}
        aliases = {}
        for line in map_file.read_text().splitlines():
            if not line.strip() or line.startswith("#"):
                continue
            p = [x.strip() for x in line.split("|")] + ["", "", "", ""]
            if p[1].startswith("="):
                aliases[p[0]] = p[1][1:]
                continue
            food = None if p[1] == "-" else self.by_name[p[1]]  # KeyError = a typo in the map: fail the build
            extra = {a.strip() for a in p[4].split(",") if a.strip()}
            self.map[p[0]] = (food, num(p[2]), num(p[3]), extra)
            if food is not None:
                food["allergens"] |= extra  # true of the food itself (pesto has pine nuts), not just this recipe
        for a, target in aliases.items():
            self.map[a] = self.map[target]  # KeyError = alias to a phrase that has no line
        self.pool = [f for f in foods if f["generic"] and f["src"] in ("usda_sr28", "usda_sr", "usda_fdc_foundation")
                     and not re.search(r"^(babyfood|infant formula|fast foods?|restaurant|school lunch)|alaska native|"
                                       r"\(navajo|\(hopi|\(apache|\(shoshone|northern plains|\brestaurant\b", f["name"], re.I)]
        self.index: dict[str, list[int]] = defaultdict(list)
        self.ptoks = []
        for i, f in enumerate(self.pool):
            t = self._toks(f["name"])
            self.ptoks.append((set(t), set(self._toks(f["name"].split(",")[0]))))
            for w in set(t):
                self.index[w].append(i)
        self.cache: dict[str, tuple | None] = {}
        self.auto_used: Counter = Counter()
        self.siblings: dict[str, list[dict]] = defaultdict(list)
        for f in self.pool:
            self.siblings[f["name"].split(",")[0]].append(f)
        self.dcache: dict[int, float | None] = {}

    @staticmethod
    def _toks(s: str) -> list[str]:
        return [w[:-1] if w.endswith("s") and len(w) > 3 else w for w in re.findall(r"[a-z]+", s.lower())]

    def lookup(self, name: str) -> tuple | None:
        if name in self.cache:
            return self.cache[name]
        hit = self.map.get(name)
        if hit is None:
            for alt in (name + "s", name[:-1] if name.endswith("s") else None, name[:-2] if name.endswith("es") else None):
                if alt and alt in self.map:
                    hit = self.map[alt]
                    break
        if hit is None:
            hit = self._auto(name)
        self.cache[name] = hit
        return hit

    def _auto(self, name: str) -> tuple | None:
        """Tail of the long tail: every word of the ingredient must be in the food's name, the head
        noun in its first part. Recipes using these must still pass the calorie cross-check."""
        pt = [w for w in self._toks(name) if w not in CUT_WORDS] or self._toks(name)
        if not pt or len(pt) > 5:
            return None
        head = pt[-1]
        best = None
        for i in self.index.get(head, ()):
            ft, first = self.ptoks[i]
            if not set(pt) <= ft or not (set(pt) & first):
                continue
            score = (len(ft - set(pt)) - ("raw" in ft) * 1.5 - (set(pt) <= first) * 2, len(self.pool[i]["name"]))
            if best is None or score < best[0]:
                best = (score, i)
        if best is None:
            return None
        self.auto_used[name] += 1
        return (self.pool[best[1]], None, None, set(), "auto")

    @staticmethod
    def g_per_ml(food: dict, cup_g: float | None) -> float | None:
        if cup_g:
            return cup_g / 236.6
        for desc, g in food["weights"]:
            if m := VOLUME_DESC.match(desc):
                amount = ing.qty_value(m[1]) if m[1] else 1.0
                unit = m[2]
                ml = ing.unit_kind("fl oz" if unit == "fl oz" else unit)[1]
                if amount and ml:
                    return g / (amount * ml)
        for rx, d in LIQUID:
            if rx.search(food["name"]):
                return d
        return None

    def density(self, food: dict, cup_g: float | None) -> float | None:
        """Own household weights, else a sibling's ('Onions, red, raw' borrows from 'Onions, raw')."""
        if (d := self.g_per_ml(food, cup_g)) is not None:
            return d
        if food["id"] not in self.dcache:
            ft = set(self._toks(food["name"]))
            first = food["name"].split(",")[0]
            sibs = [f for f in self.siblings.get(first, ()) if f is not food and self.g_per_ml(f, None)]
            best = max(sibs, key=lambda f: (len(ft & set(self._toks(f["name"]))), -len(f["name"])), default=None)
            self.dcache[food["id"]] = self.g_per_ml(best, None) if best else None
        return self.dcache[food["id"]]

    @staticmethod
    def piece_g(food: dict, piece: float | None, unit: str | None) -> float | None:
        size = unit if unit in ("large", "small", "medium") else None
        counts = [(d, g) for d, g in food["weights"] if not VOLUME_DESC.match(d) and not re.search(r"\b(oz|lb|g)\b", d)]
        if unit in ("slice", "slices"):
            for d, g in counts:
                if "slice" in d:
                    return g / (ing.qty_value(d.split()[0]) or 1)
        if size:
            for d, g in counts:
                if size in d:
                    return g / (ing.qty_value(d.split()[0]) if re.match(r"[\d.]", d) else 1)
        if piece:
            return piece * SIZE.get(size, 1.0)
        for d, g in counts:
            if re.match(r"^[\d.]+ (medium|whole|fruit|item|piece|egg|clove|pepper|stalk|potato|onion|carrot|banana|apple)\b", d):
                return g / (ing.qty_value(d.split()[0]) or 1) * SIZE.get(size, 1.0)
        return None

    def grams(self, qty: float | None, unit: str | None, hit: tuple) -> float | None:
        food, piece, cup_g, *_ = hit
        if qty is None:
            return 1.0  # 'salt to taste', 'oil for frying': a trace, still counted for allergens
        kind, factor = ing.unit_kind(unit)
        if kind == "mass":
            # Weights in recipes are as bought: '3 lbs chicken' includes the bones (USDA refuse %).
            return qty * factor * (1 - food.get("refuse", 0) / 100)
        if kind == "volume":
            d = self.density(food, cup_g)
            return qty * factor * d if d else None
        u = (unit or "").lower()
        if u == "dozen":
            u, qty = "", qty * 12
        if u in ("stick", "sticks") and re.match(r"^(butter|margarine)", food["name"], re.I):
            return qty * 113
        if u in ("can", "cans", "tin", "tins", "jar", "jars", "package", "packages", "pkg", "pkgs", "box", "boxes",
                 "bag", "bags", "carton", "cartons", "bottle", "bottles", "envelope", "envelopes", "packet", "packets",
                 "pack", "packs", "bunch", "bunches", "head", "heads", "loaf", "loaves", "ball", "balls"):
            return qty * piece if piece else None  # only foods whose map line gives the standard size
        if u in ("sprig", "sprigs"):
            return qty * 1.0
        if u in ("handful", "handfuls"):
            return qty * 30.0
        if u in ("knob", "knobs"):
            return qty * 15.0
        if u in ("leaf", "leaves") and piece:
            return qty * piece
        p = self.piece_g(food, piece, u or None)
        return qty * p if p else None


def chosen_text(rest: str) -> str:
    """'butter or margarine (or oil)' -> 'butter': the alternative the recipe was matched on."""
    t = re.sub(r"\([^)]*\)?", " ", rest)
    t = re.split(r"\bor\b|/|\band/or\b", t)[0]
    return re.split(r",| - |\bfor\b", t)[0]  # 'sugar, for dusting the cake': the use is not the food


def map_line(m: Matcher, line: str) -> tuple[str, dict | None, float, set, bool] | str:
    """-> (name, food or None for 'ignore', grams, extra allergens, auto) or a failure reason."""
    qty, unit, rest = ing.parse_line(line)
    name = ing.normalize_name(rest)
    if not name:
        return "empty ingredient line" if not rest else f"unmatched: {rest[:40]}"
    if ing.NEGLIGIBLE.match(name):
        return (name, None, 0.0, set(), False)
    # '3 cups fusilli, cooked': the state after the comma changes the food (dry pasta is 3x the kcal).
    hit = None
    if re.search(r"\bcooked\b", rest, re.I) and not name.startswith("cooked "):
        hit = m.map.get("cooked " + name)
    hit = hit or m.lookup(name)
    if hit is None:
        return f"unmatched: {name}"
    food = hit[0]
    if food is None:
        return (name, None, 0.0, set(), False)
    # The line's own words must not imply an allergen the matched food lacks ('ritz crackers' -> gluten):
    # a wrong match must never hide an allergen from the filter.
    if missing := allergens_from_text(chosen_text(rest)) - food["allergens"]:
        return f"allergen mismatch: {name} ({', '.join(sorted(missing))})"
    g = m.grams(qty, unit, hit)
    if g is None or g <= 0:
        return f"no grams: {unit or 'piece'} of {name}"
    if g > 5000:
        return "implausible amount"
    return (name, food, g, hit[3], len(hit) > 4)


# --------------------------------------------------------------------------- recipes

def lst(v) -> list:
    """JSON list, Python-literal list, R c("..") vector, or already a list."""
    if v is None:
        return []
    if isinstance(v, list):
        return v
    s = str(v).strip()
    if not s or s in ("NA", "character(0)"):
        return []
    if s.startswith("c("):
        return re.findall(r'"((?:[^"\\]|\\.)*)"', s)
    for parse in (json.loads, ast.literal_eval):
        try:
            out = parse(s)
            return list(out) if isinstance(out, (list, tuple, set)) else [out]
        except Exception:  # noqa: BLE001 - try the next format
            pass
    if m := re.fullmatch(r"""\[\s*(["'])(.*)\1\s*\]""", s, re.S):  # unescaped quotes inside: KELLOGG"S, 9"
        q = m[1]
        return [x.replace("\\" + q, q) for x in re.split(rf"{q}\s*,\s*{q}", m[2])]
    return []


def iso_minutes(s) -> int | None:
    """'PT1H30M' -> 90; '1 hrs 30 mins' -> 90; '25 mins - 30 mins' -> 25; '180' -> 180."""
    if s is None:
        return None
    s = str(s)
    if m := re.fullmatch(r"P(?:\d+D)?T?(?:(\d+)H)?(?:(\d+)M)?(?:\d+S)?", s):
        mins = int(m[1] or 0) * 60 + int(m[2] or 0)
        return mins or None
    h = re.search(r"(\d+)\s*(?:hrs?|hours?|h\b)", s)
    mm = re.search(r"(\d+)\s*(?:mins?|minutes?|m\b)", s)
    if h or mm:
        return (int(h[1]) * 60 if h else 0) + (int(mm[1]) if mm else 0) or None
    return int(float(s)) if re.fullmatch(r"\d+(\.\d+)?", s) else None


TAG_MAP = {  # source tag (lowercase) -> app tag
    "breakfast": "breakfast", "brunch": "breakfast", "lunch": "lunch", "main-dish": "dinner", "main dish": "dinner",
    "dinner": "dinner", "dinner-party": "dinner", "main course": "dinner", "one-dish-meal": "dinner",
    "snacks": "snack", "snack": "snack", "appetizers": "snack", "appetizer": "snack", "desserts": "dessert",
    "dessert": "dessert", "vegan": "vegan", "vegetarian": "vegetarian", "low-carb": "low-carb",
    "very-low-carbs": "low-carb", "low carb": "low-carb", "low-fat": "low-fat", "low-sodium": "low-sodium",
    "high-protein": "high-protein", "high protein": "high-protein", "gluten-free": "gluten-free",
    "gluten free": "gluten-free", "diabetic": "diabetic", "healthy": "healthy", "healthy-2": "healthy",
    "low-calorie": "low-calorie", "low-cholesterol": "low-cholesterol", "high-fiber": "high-fiber",
    "high in...": None, "soups-stews": "soup", "soups": "soup", "soup": "soup", "stews": "soup", "salads": "salad",
    "salad": "salad", "beverages": "drink", "smoothies": "drink", "drinks": "drink", "kid-friendly": "kid-friendly",
    "easy": "easy", "15-minutes-or-less": "quick", "30-minutes-or-less": "quick", "< 15 mins": "quick",
    "< 30 mins": "quick", "baking": "baking", "breads": "baking", "cakes": "baking", "cookies-and-brownies": "baking",
    **{c: c for c in ("italian", "mexican", "indian", "chinese", "japanese", "thai", "greek", "french", "spanish",
                      "korean", "vietnamese", "moroccan", "middle-eastern", "mediterranean", "caribbean", "german",
                      "british", "american", "asian", "african", "cajun", "southern", "tex-mex")},
}
SOFT = re.compile(r"\b(soup|porridge|mash|mashed|custard|pudding|oatmeal|congee|risotto|scrambled|mousse|"
                  r"polenta|grits|stew|dal|dhal|hummus|yogh?urt|jelly|jello)\b", re.I)
PUREED = re.compile(r"\b(pur[eé]e[ds]?|smoothie|baby food|blended)\b", re.I)


def tags_of(raw_tags: list[str], minutes: int | None) -> str:
    out = []
    for t in raw_tags:
        a = TAG_MAP.get(str(t).strip().lower())
        if a and a not in out:
            out.append(a)
    if minutes is not None and minutes <= 30 and "quick" not in out:
        out.append("quick")
    return ",".join(out)


def texture_of(title: str, tags: str) -> str | None:
    # ponytail: keyword heuristic on the title; a curated texture list if ALS users report misses
    if PUREED.search(title):
        return "pureed"
    if SOFT.search(title) or "soup" in tags.split(","):
        return "soft"
    return None


def instructions_of(steps) -> str:
    if isinstance(steps, str) and not steps.strip().startswith(("[", "c(")):
        return re.sub(r"\s+\n", "\n", html.unescape(steps).strip())
    parts = [re.sub(r"\s+", " ", html.unescape(str(s))).strip() for s in lst(steps)]
    parts = [p for p in parts if p]
    return "\n".join(f"{i}. {p}" for i, p in enumerate(parts, 1)) if len(parts) > 1 else (parts[0] if parts else "")


class Recipes:
    """Collects checked recipes from every source; dedupes on title + foods."""

    def __init__(self, m: Matcher):
        self.m = m
        self.m_foods = {f["id"]: f for f in m.by_name.values()}
        self.kept: dict[tuple, dict] = {}
        self.unmatched: Counter = Counter()

    def add(self, rep: dict, src: str, sid, title, lines, steps, servings, minutes=None, raw_tags=(),
            published_kcal: float | None = None, rank: float = 0.0, curated_only=False) -> None:
        title = tidy(html.unescape(str(title or "")))
        instructions = instructions_of(steps)
        try:
            servings = int(round(float(servings))) if servings not in (None, "") else None
        except (TypeError, ValueError):
            servings = None
        if not title:
            return self._drop(rep, "no title")
        if not lines:
            return self._drop(rep, "no ingredient lines")
        if not instructions:
            return self._drop(rep, "no instructions")
        if not servings or servings < 1 or servings > 60:
            return self._drop(rep, "no servings")
        items: dict[int, float] = {}
        names: dict[int, str] = {}  # food id -> the line's own short name ('walnuts'), for retitling
        auto = False
        extra_allergens: dict[int, set] = defaultdict(set)
        for line in lines:
            if not str(line).strip():
                continue
            got = map_line(self.m, html.unescape(str(line)))
            if isinstance(got, str):
                if got.startswith(("unmatched: ", "no grams")):
                    self.unmatched[got] += 1
                if got.startswith("allergen mismatch"):
                    self.unmatched[got] += 1
                    return self._drop(rep, "ingredient words imply an allergen the matched food lacks")
                return self._drop(rep, "ingredient not matched to a food" if got.startswith("unmatched") else
                                  "ingredient amount not convertible to grams" if got.startswith("no grams") else got)
            name, food, g, extra, was_auto = got
            if food is None:
                continue
            items[food["id"]] = items.get(food["id"], 0) + g
            names.setdefault(food["id"], name)
            extra_allergens[food["id"]] |= extra
            auto |= was_auto
        if not items:
            return self._drop(rep, "no ingredient with nutrition")
        if auto and curated_only:
            return self._drop(rep, "ingredient matched only by name similarity, no published calories to confirm")
        foods = self.m_foods
        kcal = sum(foods[f]["kcal"] * g / 100 for f, g in items.items()) / servings
        if not 15 <= kcal <= 2500:
            return self._drop(rep, "calories per serving implausible")
        if published_kcal and published_kcal > 0:
            if abs(kcal - published_kcal) > max(80, 0.3 * published_kcal):
                return self._drop(rep, "calories differ from the source's published value")
            rank += 1  # confirmed by the source
        minutes = minutes if minutes and 0 < minutes <= 60 * 24 * 3 else None
        tags = tags_of(list(raw_tags), minutes)
        k = (key(title), frozenset(items))
        if k in self.kept:
            return self._drop(rep, "duplicate of a recipe already kept")
        self.kept[k] = {"id": stable_id("recipe", src, sid), "title": title[:200], "instructions": instructions,
                        "servings": servings, "minutes": minutes, "tags": tags, "texture": texture_of(title, tags),
                        "items": items, "names": names, "rank": rank, "src": src, "extra_allergens": extra_allergens, "file": rep}
        rep["kept"] += 1

    @staticmethod
    def _drop(rep: dict, why: str) -> None:
        rep["dropped"][why] += 1


def load_foodcom(raw: Path, R: Recipes) -> None:
    import duckdb
    rr = raw / "raw_recipes"
    con = duckdb.connect()
    con.execute("SET memory_limit='1500MB'")
    meta = {}
    rep_p = report("raw_recipes/recipes.parquet", "recipe metadata (food.com)",
                   "food.com: total time, keywords, category, rating, published calories per serving. Ingredient "
                   "amounts here have no units, so lines come from recipes_ingredients.csv / recipes_w_search_terms.csv.")
    cur = con.execute("SELECT RecipeId::BIGINT, TotalTime, Keywords, RecipeCategory, AggregatedRating, ReviewCount, "
                      "Calories, RecipeServings FROM read_parquet(?)", [str(rr / "recipes.parquet")])
    while batch := cur.fetchmany(20000):
        for r in batch:
            rep_p["rows"] += 1
            meta[r[0]] = (iso_minutes(r[1]), list(r[2] or []) + ([r[3]] if r[3] else []), r[4], r[5], r[6], r[7])
    rep_c = report("raw_recipes/recipes.csv", "same as recipes.parquet",
                   "CSV export of recipes.parquet: same 522,517 RecipeIds (checked), not read twice.")
    rep_c["rows"] = rep_p["rows"]
    rep_c["dropped"]["duplicate of recipes.parquet"] = rep_p["rows"]
    raw_meta = {}
    rep_r = report("raw_recipes/RAW_recipes.csv", "recipe metadata (food.com)",
                   "food.com: minutes, tags, per-serving calories. Ingredients without amounts.")
    for r in rows(rr / "RAW_recipes.csv"):
        rep_r["rows"] += 1
        nut = lst(r["nutrition"])
        raw_meta[int(r["id"])] = (num(r["minutes"]), lst(r["tags"]), num(nut[0]) if nut else None)
    seen = set()
    for file, lines_col in (("recipes_ingredients.csv", "ingredients_raw"), ("recipes_w_search_terms.csv", "ingredients_raw_str")):
        rep = report("raw_recipes/" + file, "recipes (food.com)",
                     "food.com recipes with full ingredient lines, servings and serving weight.")
        for r in rows(rr / file):
            rep["rows"] += 1
            try:
                rid = int(r["id"])
            except ValueError:
                rep["dropped"]["no id"] += 1
                continue
            if rid in seen:
                rep["dropped"]["same food.com recipe as recipes_ingredients.csv"] += 1
                continue
            seen.add(rid)
            mt = meta.get(rid)
            rw = raw_meta.get(rid)
            minutes = (mt and mt[0]) or (rw and rw[0] and int(rw[0])) or None
            tags = lst(r["tags"]) + (mt[1] if mt else []) + (rw[1] if rw else []) + lst(r.get("search_terms"))
            kcal = (mt and mt[4]) or (rw and rw[2]) or None
            rating = (mt[2] or 0) * math.log1p(mt[3] or 0) if mt else 0
            R.add(rep, "food.com", rid, r["name"], lst(r[lines_col]), r["steps"], num(r["servings"]) or (mt and mt[5]),
                  minutes, tags, kcal, rank=rating)
    rep_p["kept"] = sum(1 for rid in meta if rid in seen)
    rep_p["dropped"]["no food.com row with ingredient amounts"] = rep_p["rows"] - rep_p["kept"]
    rep_r["kept"] = sum(1 for rid in raw_meta if rid in seen)
    rep_r["dropped"]["no food.com row with ingredient amounts"] = rep_r["rows"] - rep_r["kept"]
    rep_r["note"] += f" {rep_r['kept']} rows give minutes/tags/calories to recipes kept from the files with amounts."
    rep_p["note"] += f" {rep_p['kept']} rows give metadata to recipes read from the files with amounts."


def load_bbc(raw: Path, R: Recipes) -> None:
    for file in ("baking.json", "health.json", "recipes_97.json"):
        rep = report("raw_recipes/" + file, "recipes (BBC Good Food)", "Metric ingredient lines, serves, published kcal per serving.")
        for r in json.load(open(raw / "raw_recipes" / file)):
            rep["rows"] += 1
            t = r.get("times") or {}
            minutes = sum(filter(None, (iso_minutes(t.get("Preparation")), iso_minutes(t.get("Cooking"))))) or None
            kcal = num((r.get("nutrients") or {}).get("kcal"))
            rank = (r.get("rattings") or 0) * math.log1p(r.get("vote_count") or 0)
            R.add(rep, "bbcgoodfood", r.get("url") or r.get("id"), r.get("name"), r.get("ingredients") or [],
                  r.get("steps") or [], num(r.get("serves")), minutes,
                  [r.get("subcategory") or "", r.get("dish_type") or "", r.get("maincategory") or ""], kcal, rank=rank)


def load_small_recipes(raw: Path, R: Recipes) -> None:
    rr = raw / "raw_recipes"
    rep = report("raw_recipes/recipes_99.csv", "recipes (allrecipes)", "Servings and per-serving fat/carbs/protein (kcal derived).")
    for r in rows(rr / "recipes_99.csv"):
        rep["rows"] += 1
        nut = r["nutrition"] or ""
        g = lambda label: num((re.search(label + r"\s+([\d.]+)\s*g", nut) or [None, None])[1])  # noqa: E731
        f, c, p = g("Total Fat"), g("Total Carbohydrate"), g("Protein")
        kcal = 9 * f + 4 * c + 4 * p if None not in (f, c, p) else None
        lines = re.split(r",\s+(?=[\d½¼¾⅓⅔⅛])", r["ingredients"] or "")
        R.add(rep, "allrecipes99", r["url"] or r["recipe_name"], r["recipe_name"], lines, r["directions"],
              num(r["servings"]), iso_minutes(r["total_time"]) or iso_minutes(r["timing"]),
              [x.strip() for x in (r["cuisine_path"] or "").strip("/").split("/")], kcal)
    rep = report("raw_recipes/christmas_recipes.csv", "recipes (BBC, Christmas)", "Servings '10 serving(s)'; no published nutrition.")
    for r in rows(rr / "christmas_recipes.csv"):
        rep["rows"] += 1
        R.add(rep, "christmas", r["Title"], r["Title"], lst(r["Ingredients"]), r["Instructions"], num(r["Servings"]),
              iso_minutes(r["Time"]), ["dinner"], None, curated_only=True)
    for file in ("Recipes.csv", "Recipes_2.csv", "Recipes_3.csv"):
        rep = report("raw_recipes/" + file, "recipes (Reddit posts)",
                     "Free-text Reddit posts; ingredients = bullet list under an 'Ingredients' heading, servings from 'Serves/Servings: N'.")
        for r in rows(rr / file):
            rep["rows"] += 1
            text = (r["comment"] or "").replace("\r", "")
            serv = re.search(r"(?:serv(?:es|ings?)|yield|makes)\s*[:\-]?\s*\**\s*(\d+)", text, re.I)
            im = re.search(r"ingredients?\s*:?\**\s*\n(.*?)(?:\n\s*\**\s*(?:#+\s*)?(?:instructions?|directions?|method|steps|preparation)\b)",
                           text, re.I | re.S)
            if not im:
                rep["dropped"]["no ingredient list found in the post"] += 1
                continue
            lines = [re.sub(r"^\s*(?:[*\-•]|\d+\.)\s*", "", x).strip() for x in im[1].splitlines()]
            lines = [x for x in lines if x and not re.match(r"^#|^\*\*.*\*\*$|^for the\b", x, re.I)]
            steps = text[im.end():].strip()
            R.add(rep, "reddit", r["title"] + r["date"], re.sub(r"\s*recipe\s*$", "", r["title"], flags=re.I), lines, steps,
                  num(serv[1]) if serv else None, None, [], None, curated_only=True)


def count_only(raw: Path, file: str, target: str, note: str, reason: str) -> None:
    """Files that can't give recipes the app can use: counted (every row read) and explained."""
    rep = report("raw_recipes/" + file, target, note)
    p = raw / "raw_recipes" / file
    if p.suffix == ".json":
        rep["rows"] = len(json.load(open(p)))
    else:
        rep["rows"] = sum(1 for _ in rows(p))
    rep["dropped"][reason] = rep["rows"]


def load_unusable_recipes(raw: Path) -> None:
    count_only(raw, "RecipeNLG_dataset.csv", "-",
               "RecipeNLG (2.2M, Recipes1M + gathered sites). Same file as full_dataset.csv (byte-identical). "
               "Licensed for non-commercial research only.", "no servings (per-serving nutrition impossible)")
    rep = report("raw_recipes/full_dataset.csv", "-", "Byte-identical copy of RecipeNLG_dataset.csv (same MD5).")
    rep["rows"] = REPORT["raw_recipes/RecipeNLG_dataset.csv"]["rows"]
    rep["dropped"]["duplicate file (same MD5 as RecipeNLG_dataset.csv)"] = rep["rows"]
    count_only(raw, "1_Recipe_csv.csv", "-", "allrecipes.com, 62k: ingredient lines and directions.",
               "no servings (per-serving nutrition impossible)")
    count_only(raw, "recipes_extended.csv", "-", "Same 62k allrecipes rows as 1_Recipe_csv.csv plus derived labels.",
               "no servings (per-serving nutrition impossible)")
    count_only(raw, "Multi_Cuisine_Recipe_Dataset.csv", "-", "620 Indian/world recipes with amounts.",
               "no servings (per-serving nutrition impossible)")
    count_only(raw, "food_recipes.csv", "-", "archanaskitchen.com, 8k: ingredient names separated by '|'.",
               "ingredients have no amounts")
    count_only(raw, "Receipes from around the world.csv", "-", "153 recipes: ingredient names, servings, calories.",
               "ingredients have no amounts")
    count_only(raw, "recipes.json", "-", "chefkoch.de, 12k recipes in German.",
               "German text and no servings (the app is English-only for now)")


# --------------------------------------------------------------------------- write

def r3(v):
    return None if v is None else round(v, 3)


def carry_over_dev_tables(db: sqlite3.Connection, m: Matcher) -> list[str]:
    """conditions / condition_rules / docs from db/nutri_dummy.sql (no source data for them in the raw
    folders); condition_avoid re-linked by food name. Returns notes for the report."""
    dev = sqlite3.connect(":memory:")
    dev.executescript((ROOT / "db/nutri_schema.sql").read_text())
    dev.executescript((ROOT / "db/nutri_dummy.sql").read_text())
    for t in ("conditions", "condition_rules", "docs"):
        cols = [r[1] for r in dev.execute(f"PRAGMA table_info({t})")]
        db.executemany(f"INSERT INTO {t} ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                       dev.execute(f"SELECT {','.join(cols)} FROM {t}").fetchall())
    notes = []
    for cid, fname, note in dev.execute("SELECT a.condition_id, f.name, a.note FROM condition_avoid a JOIN foods f ON f.id = a.food_id"):
        hit = m.lookup(ing.normalize_name(fname.split(",")[0]))
        if hit and hit[0]:
            db.execute("INSERT OR IGNORE INTO condition_avoid VALUES (?, ?, ?)", (cid, hit[0]["id"], note))
            notes.append(f"condition_avoid: '{fname}' -> '{hit[0]['name']}'")
        else:
            notes.append(f"condition_avoid: '{fname}' has no match, left out")
    return notes


def write_db(out: Path, foods, dishes, products, recipes, m: Matcher) -> list[str]:
    tmp = out.with_suffix(".tmp")
    tmp.unlink(missing_ok=True)
    db = sqlite3.connect(tmp)
    db.execute("PRAGMA foreign_keys = ON")
    db.execute("PRAGMA journal_mode = OFF")
    db.executescript((ROOT / "db/nutri_schema.sql").read_text())
    db.executemany("INSERT INTO meta VALUES (?, ?)", [
        ("schema_version", "1"), ("embedding_model", "embeddinggemma-300m"), ("built_at", date.today().isoformat()),
        ("source", "tool/build_real_nutri_db.py; per-file report in docs/DATA_SOURCES.md")])
    db.executemany("INSERT INTO foods VALUES (?,?,?,?,?,?,?,?,?,?)", [
        (f["id"], f["name"], *(r3(f.get(k)) for k in NUTRIENTS), r3(f["serving_g"])) for f in foods])
    db.executemany("INSERT INTO allergens VALUES (?, ?)", [(f["id"], a) for f in foods for a in sorted(f["allergens"])])
    db.executemany("INSERT INTO recipes VALUES (?,?,?,?,?,?,?)", [
        (r["id"], r["title"], r["instructions"], r["servings"], r["minutes"], r["tags"], r["texture"]) for r in recipes])
    db.executemany("INSERT INTO recipe_ingredients VALUES (?,?,?)", [
        (r["id"], fid, round(g, 1)) for r in recipes for fid, g in r["items"].items() if round(g, 1) > 0])
    seen = set()
    rows_d = []
    for d in dishes:
        if key(d["name"]) in seen:
            continue
        seen.add(key(d["name"]))
        rows_d.append((stable_id("dish", d["src"], d["skey"]), d["name"], r3(d["serving_g"]), None, *(r3(d.get(k)) for k in NUTRIENTS)))
    db.executemany("INSERT INTO dishes VALUES (?,?,?,?,?,?,?,?,?,?,?)", rows_d)
    db.executemany("INSERT INTO products VALUES (?,?,?,?,?,?,?,?,?)", [
        (p["barcode"], p["name"], *(r3(p.get(k)) for k in NUTRIENTS)) for p in products])
    notes = carry_over_dev_tables(db, m)
    finish(db, tmp, out)
    return notes


def pick_recipes(recipes: list[dict], budget_bytes: int) -> list[dict]:
    """Best first (rating x reviews, +1 when the source's calories confirmed ours) until the estimated
    size of recipe text, ingredient rows and their indexes reaches the budget."""
    ranked = sorted(recipes, key=lambda r: (-r["rank"], r["id"]))
    out, used = [], 0
    for r in ranked:
        # ponytail: estimate from measured builds (text ~2.3x with FTS + page overhead, ~40 B per ingredient row)
        cost = 2.3 * (len(r["title"]) + len(r["instructions"]) + len(r["tags"]) + 40) + 40 * len(r["items"])
        if used + cost > budget_bytes:
            r["file"]["dropped"]["over the 300 MB size budget (lower-ranked)"] += 1
            r["file"]["kept"] -= 1
            continue
        used += cost
        out.append(r)
    return out


# Too common to tell two versions of a recipe apart.
STAPLES = re.compile(r"\b(salt|pepper|water|sugar|flour|baking (powder|soda)|soda|cornstarch|corn starch|yeast|"
                     r"vanilla|oil|shortening|cooking spray|spray|ground (cinnamon|nutmeg|cloves|ginger|allspice)|"
                     r"cinnamon|nutmeg|eggs?|egg whites?|egg yolks?|ice|garlic clove|garlic cloves)\b", re.I)


def retitle_duplicates(recipes: list[dict]) -> int:
    """Same title, different ingredients (40 'Banana Bread's): the best-ranked keeps the title, the
    others get what sets them apart, 'Banana Bread with walnuts', then a second ingredient, then a
    number. Ids don't change. Returns how many were renamed."""
    groups = defaultdict(list)
    for r in sorted(recipes, key=lambda r: (-r["rank"], r["id"])):
        groups[key(r["title"])].append(r)
    renamed = 0
    for group in groups.values():
        if len(group) < 2:
            continue
        used = {key(group[0]["title"])}
        count = Counter(fid for r in group for fid in r["items"])
        for r in group[1:]:
            def stems(s: str) -> set[str]:
                return {w[:-1] if len(w) > 3 and w.endswith("s") else w for w in key(s).split()}
            words = stems(r["title"])
            # One no other version has first, then the biggest amount; never a staple ("with salt")
            # or a word already in the title.
            extras = sorted((fid for fid in r["items"] if r["names"].get(fid)
                             and not STAPLES.search(r["names"][fid]) and not stems(r["names"][fid]) <= words),
                            key=lambda fid: (count[fid] > 1, -r["items"][fid]))
            names = list(dict.fromkeys(r["names"][fid] for fid in extras))
            joint = " and " if " with " in r["title"].lower() else " with "
            for title in ([f"{r['title']}{joint}{names[0]}"] if names else []) + \
                    ([f"{r['title']}{joint}{names[0]} and {names[1]}"] if len(names) > 1 else []) + \
                    [f"{r['title']} ({n})" for n in range(2, len(group) + 2)]:
                if key(title) not in used:
                    break
            used.add(key(title))
            r["title"] = title[:200]
            renamed += 1
    return renamed


def report_copies(raw: Path) -> None:
    """Every file gets a report row: an unread file must be a byte-identical copy of one that was read."""
    files = [p for d in ("raw_nutritions", "raw_recipes") for p in sorted((raw / d).rglob("*")) if p.is_file()]
    md5 = lambda p: hashlib.file_digest(open(p, "rb"), "md5").hexdigest()  # noqa: E731
    for p in files:
        rel = str(p.relative_to(raw))
        if rel in REPORT:
            continue
        twin = next((str(q.relative_to(raw)) for q in files if str(q.relative_to(raw)) in REPORT
                     and q.stat().st_size == p.stat().st_size and md5(q) == md5(p)), None)
        if twin is None:
            raise SystemExit(f"{rel} is not read by any loader and is not a copy of a file that is")
        rep = report(rel, "-", f"Byte-identical copy of `{twin}` (same MD5), read once.")
        rep["rows"] = REPORT[twin]["rows"]
        rep["dropped"][f"duplicate file (same MD5 as {Path(twin).name})"] = rep["rows"]


def write_report(path: Path, counts: dict, notes: list[str]) -> None:
    lines = ["# nutri.db data sources", "",
             f"Built by `tool/build_real_nutri_db.py` on {date.today().isoformat()} from "
             "`nutrition-ai-assistant/data/raw_nutritions` and `raw_recipes`. Every file in both folders is listed; "
             "byte-identical copies are read once.", "",
             "| table | rows |", "|---|---|", *[f"| {k} | {v:,} |" for k, v in counts.items()], "",
             "## Rules", "",
             "- Non-authoritative nutrition rows must pass range checks and kcal ≈ 4·protein + 4·carbs + 9·fat "
             "(±20 %, ±20 kcal; alcohol exempt).",
             "- A recipe is kept only if it states servings and **every** ingredient line maps to a food and a gram "
             "amount (water, salt-and-pepper-to-taste and similar are allowed to carry no nutrition).",
             "- An ingredient line whose words imply an allergen the matched food lacks is rejected, so a wrong "
             "match can never hide an allergen from the filter.",
             "- Where the source publishes calories per serving, ours must agree within ±30 % (±80 kcal).",
             "- Recipes are ranked by rating × log(reviews) (+1 when calories were confirmed) and kept until the "
             "300 MB budget.", "",
             "## Per file", "",
             "| file | rows | kept | used for | dropped (reason: rows) | notes |", "|---|---|---|---|---|---|"]
    for f in sorted(REPORT):
        r = REPORT[f]
        dropped = "; ".join(f"{k}: {v:,}" for k, v in r["dropped"].most_common() if v)
        lines.append(f"| `{f}` | {r['rows']:,} | {max(r['kept'], 0):,} | {r['target']} | {dropped or '-'} | {r['note']} |")
    lines += ["", "## Carried over from the development data", "",
              "No raw file covers conditions, condition rules or medical docs (`raw_medical/` was not part of this "
              "build), so those rows are copied unchanged from `db/nutri_dummy.sql` and are still **development "
              "values**:", "", *[f"- {n}" for n in notes], "",
              "## Licensing (data owner to confirm before release)", "",
              "- USDA SR28 / FoodData Central: public domain (US government).",
              "- NZ FOODfiles: © Plant & Food Research / Ministry of Health; check the FOODfiles licence terms.",
              "- Open Food Facts: Open Database License (ODbL) — attribution and share-alike of the database.",
              "- food.com, BBC Good Food, allrecipes, Reddit text: scraped content; copyright stays with the "
              "authors/sites. RecipeNLG (not used) is non-commercial research only.", ""]
    path.write_text("\n".join(lines))


def main(raw: Path, out: Path) -> None:
    sr = load_sr(raw)
    sr += load_food_csv(raw, sr)
    fdc, nz = load_fdc(raw), load_nz(raw)
    flat, dishes = load_flat_foods(raw)
    foods = merge_foods([sr, fdc, nz, flat])
    m = Matcher(foods, Path(__file__).parent / "ingredient_map.txt")
    R = Recipes(m)
    load_foodcom(raw, R)
    load_bbc(raw, R)
    load_small_recipes(raw, R)
    load_unusable_recipes(raw)
    products = load_off_products(raw)
    # Budget, in pick_recipes' estimate: it runs ~1.35x the real bytes (2026-10-06: 170 MB estimated = 126 MB
    # of recipes + FTS). 300 MB estimated ~ 222 MB real, plus ~62 MB other tables, stays under the contract.
    budget = int(sys.argv[3]) if len(sys.argv) > 3 else 300_000_000
    recipes = pick_recipes(list(R.kept.values()), budget)
    renamed = retitle_duplicates(recipes)
    used_foods = {fid for r in recipes for fid in r["items"]}
    notes = write_db(out, foods, dishes, products, recipes, m)
    size = out.stat().st_size
    db = sqlite3.connect(out)
    counts = {t: db.execute(f"SELECT count(*) FROM {t}").fetchone()[0]
              for t in ("foods", "allergens", "recipes", "recipe_ingredients", "dishes", "products", "conditions",
                        "condition_rules", "condition_avoid", "docs")}
    db.close()
    counts["nutri.db bytes"] = size
    counts["foods used by recipes"] = len(used_foods)
    counts["recipes renamed (same title, other ingredients)"] = renamed
    report_copies(raw)
    write_report(ROOT / "docs/DATA_SOURCES.md", counts, notes)
    print(json.dumps(counts, indent=1))
    # What to add to ingredient_map.txt next: each line here blocked that many recipes.
    print("top unmatched ingredient lines:", *(f"  {n:6} {why}" for why, n in R.unmatched.most_common(30)), sep="\n")
    if size > MAX_DB_BYTES:
        raise SystemExit(f"nutri.db is {size:,} bytes, over the {MAX_DB_BYTES:,} contract: lower the budget")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]).expanduser(), Path(sys.argv[2]))
